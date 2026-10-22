"""
Memory RAG Test using QRECC Dataset - FastAPI TestClient Version

This version uses TestClient for fast integration testing without requiring
a running server. Uses the same logic as test_qrecc_memory.py but with
TestClient for direct in-process testing.
"""
import sys
import json
import os
import time
from pathlib import Path
from collections import defaultdict
from dataclasses import dataclass
from typing import List, Dict, Any, Optional

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from dotenv import load_dotenv
from fastapi.testclient import TestClient

# Add project root to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

load_dotenv()

from main import app, INTERNAL_API_SECRET

client = TestClient(app)

HEADERS = {
    "X-Internal-Secret": INTERNAL_API_SECRET,
    "Content-Type": "application/json"
}

QRECC_FILE = Path(__file__).resolve().parent / "data" / "memory_dataset.csv"
UPLOAD_URL = "/api/v1/documents/upload"
CHAT_URL = "/api/v1/chat/query"
CONV_DETAIL_URL = "/api/v1/chat/conversations"


@dataclass
class QrecTurn:
    conversation_no: int
    turn_no: int
    question: str
    rewrite: str
    answer: str
    context: List[str]


@dataclass
class TestResult:
    conversation_no: int
    turn_no: int
    question: str
    expected_rewrite: str
    actual_rewrite: str
    expected_answer: str
    actual_answer: str
    rewrite_correct: bool
    answer_correct: bool
    conversation_id: str


def load_qrecc_data() -> List[QrecTurn]:
    """Load and parse ViWiki-QR CSV test data."""
    import pandas as pd
    df = pd.read_csv(QRECC_FILE)
    
    turns = []
    for _, row in df.iterrows():
        turns.append(QrecTurn(
            conversation_no=int(row["session_id"]),
            turn_no=int(row["turn_id"]) + 1,
            question=str(row["question"]),
            rewrite=str(row["rewrite"]),
            answer=str(row["answer"]),
            context=[]  # Not in dataset
        ))
    return turns


def group_by_conversation(turns: List[QrecTurn]) -> Dict[int, List[QrecTurn]]:
    """Group turns by conversation number."""
    grouped = defaultdict(list)
    for turn in turns:
        grouped[turn.conversation_no].append(turn)
    for conv_no in grouped:
        grouped[conv_no].sort(key=lambda t: t.turn_no)
    return grouped


def select_conversations(grouped: Dict[int, List[QrecTurn]], min_turns: int = 2, max_convs: int = None) -> List[int]:
    """Select all conversations with at least min_turns (default: all 250)."""
    candidates = [(conv_no, len(turns)) for conv_no, turns in grouped.items() if len(turns) >= min_turns]
    candidates.sort(key=lambda x: (-x[1], x[0]))
    selected = [conv_no for conv_no, _ in candidates] if max_convs is None else [conv_no for conv_no, _ in candidates[:max_convs]]
    print(f"Selected {len(selected)} conversations (min {min_turns} turns)")
    return selected


def create_docx_for_conversation(turns: List[QrecTurn], output_path: Path) -> None:
    """Create a .docx file with Rewrite as heading and Answer as body for each turn."""
    doc = docx.Document()
    
    # Add title
    doc.add_heading(f"QRECC Conversation {turns[0].conversation_no} - Knowledge Base", level=0)
    
    for turn in turns:
        # Rewrite as heading
        doc.add_heading(f"Turn {turn.turn_no}: {turn.rewrite}", level=1)
        # Answer as paragraph
        doc.add_paragraph(turn.answer)
        doc.add_paragraph("")  # spacing
    
    doc.save(output_path)
    print(f"Created {output_path} with {len(turns)} turns")


def upload_document(file_path: Path) -> Optional[str]:
    """Upload a .docx file and return document_id."""
    with open(file_path, "rb") as f:
        files = {"file": (file_path.name, f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        data = {
            "allowed_roles": "ROLE_EMPLOYEE,ROLE_MANAGER,ROLE_ADMIN",
            "allowed_departments": "ALL",
            "uploaded_by": "QRECC_TEST"
        }
        
        print(f"  Uploading {file_path.name}...")
        response = client.post(UPLOAD_URL, files=files, data=data)
        
        if response.status_code == 200:
            result = response.json()
            doc_id = result.get("document_id")
            print(f"  Uploaded: {doc_id} (status: {result.get('status')})")
            return doc_id
        else:
            print(f"  Upload failed: {response.status_code} - {response.text}")
            return None


def wait_for_processing(document_id: str, max_wait: int = 180) -> bool:
    """Wait for document processing to complete."""
    print(f"  Waiting for processing of {document_id}...")
    start = time.time()
    while time.time() - start < max_wait:
        try:
            resp = client.get(f"{CONV_DETAIL_URL.replace('/chat/conversations', '/documents')}/{document_id}", headers=HEADERS)
            if resp.status_code == 200:
                data = resp.json()
                status = data.get("status")
                if status == "COMPLETED":
                    print(f"  Processing completed")
                    return True
                elif status == "FAILED":
                    print(f"  Processing failed: {data.get('error_message', 'Unknown error')}")
                    return False
                print(f"  Status: {status}...")
        except Exception as e:
            print(f"  Error checking status: {e}")
        time.sleep(3)
    print(f"  Timeout waiting for processing")
    return False


def chat_query(question: str, conversation_id: Optional[str], user_id: str) -> Dict[str, Any]:
    """Send a chat query and return the response."""
    payload = {
        "question": question,
        "conversation_id": conversation_id,
        "user_id": user_id,
        "user_roles": ["ROLE_EMPLOYEE"],
        "department": "ALL",
        "top_k": 3
    }
    
    response = client.post(CHAT_URL, headers=HEADERS, json=payload)
    if response.status_code != 200:
        raise Exception(f"Chat query failed: {response.status_code} - {response.text}")
    return response.json()


def get_conversation_detail(conversation_id: str, user_id: str) -> Dict[str, Any]:
    """Get conversation details including context_summary."""
    resp = client.get(
        f"{CONV_DETAIL_URL}/{conversation_id}",
        headers=HEADERS,
        params={"user_id": user_id}
    )
    if resp.status_code != 200:
        raise Exception(f"Get conversation failed: {resp.status_code} - {resp.text}")
    return resp.json()


def evaluate_rewrite(expected: str, actual: str) -> bool:
    """Check if actual rewrite contains key entities from expected rewrite."""
    expected_lower = expected.lower()
    actual_lower = actual.lower()
    
    import re
    expected_words = set(re.findall(r'\b\w{4,}\b', expected_lower))
    actual_words = set(re.findall(r'\b\w{4,}\b', actual_lower))
    
    key_entities = [w for w in expected_words if len(w) > 4]
    
    if not key_entities:
        return True
    
    matches = sum(1 for e in key_entities if e in actual_lower)
    return matches / len(key_entities) >= 0.7


def evaluate_answer(expected: str, actual: str) -> bool:
    """Check if actual answer contains key information from expected answer."""
    expected_lower = expected.lower()
    actual_lower = actual.lower()
    
    import re
    expected_words = set(re.findall(r'\b\w{4,}\b', expected_lower))
    actual_words = set(re.findall(r'\b\w{4,}\b', actual_lower))
    
    key_entities = [w for w in expected_words if len(w) > 4]
    
    if not key_entities:
        return True
    
    matches = sum(1 for e in key_entities if e in actual_lower)
    return matches / len(key_entities) >= 0.6


def run_conversation_test(conv_no: int, turns: List[QrecTurn], user_id: str) -> List[TestResult]:
    """Run test for a single conversation."""
    print(f"\n{'='*60}")
    print(f"TESTING CONVERSATION {conv_no} ({len(turns)} turns)")
    print(f"{'='*60}")
    
    results = []
    conversation_id = None
    
    for i, turn in enumerate(turns):
        print(f"\n--- Turn {turn.turn_no}/{len(turns)} ---")
        print(f"Q: {turn.question}")
        print(f"Expected Rewrite: {turn.rewrite}")
        
        try:
            result = chat_query(turn.question, conversation_id, user_id)
            conversation_id = result["conversation_id"]
            
            actual_rewrite = result.get("rewritten_query", "")
            actual_answer = result.get("answer", "")
            
            rewrite_ok = evaluate_rewrite(turn.rewrite, actual_rewrite)
            answer_ok = evaluate_answer(turn.answer, actual_answer)
            
            print(f"Actual Rewrite: {actual_rewrite}")
            print(f"Rewrite Match: {'✅' if rewrite_ok else '❌'}")
            print(f"Answer Match: {'✅' if answer_ok else '❌'}")
            
            results.append(TestResult(
                conversation_no=conv_no,
                turn_no=turn.turn_no,
                question=turn.question,
                expected_rewrite=turn.rewrite,
                actual_rewrite=actual_rewrite,
                expected_answer=turn.answer,
                actual_answer=actual_answer,
                rewrite_correct=rewrite_ok,
                answer_correct=answer_ok,
                conversation_id=conversation_id
            ))
            
        except Exception as e:
            print(f"Error on turn {turn.turn_no}: {e}")
            results.append(TestResult(
                conversation_no=conv_no,
                turn_no=turn.turn_no,
                question=turn.question,
                expected_rewrite=turn.rewrite,
                actual_rewrite="ERROR",
                expected_answer=turn.answer,
                actual_answer=str(e),
                rewrite_correct=False,
                answer_correct=False,
                conversation_id=conversation_id or ""
            ))
        
        time.sleep(0.5)  # Small delay
    
    # Check summary memory retention at the end
    if conversation_id:
        print(f"\n--- Checking Summary Memory Retention ---")
        try:
            detail = get_conversation_detail(conversation_id, user_id)
            context_summary = detail.get("context_summary", "")
            message_count = detail.get("message_count", 0)
            
            print(f"Message Count: {message_count}")
            print(f"Context Summary: {context_summary[:200]}...")
            
            # Check if first turn entity is preserved in summary
            first_turn = turns[0]
            first_entity_keywords = set(first_turn.rewrite.lower().split())
            summary_lower = context_summary.lower()
            
            # Key entities from first turn
            first_key_words = [w for w in first_entity_keywords if len(w) > 4]
            retained = sum(1 for w in first_key_words if w in summary_lower)
            
            print(f"First turn key words retained in summary: {retained}/{len(first_key_words)}")
            
            # Test final turn asking about first turn
            if len(turns) >= 5:
                print(f"\n--- Final Memory Test: Ask about Turn 1 ---")
                memory_question = f"What was the first question about?"
                mem_result = chat_query(memory_question, conversation_id, user_id)
                mem_rewrite = mem_result.get("rewritten_query", "")
                mem_answer = mem_result.get("answer", "")
                print(f"Memory Rewrite: {mem_rewrite}")
                print(f"Memory Answer: {mem_answer[:200]}...")
                
                # Check if rewrite contains first turn entity
                memory_rewrite_ok = evaluate_rewrite(first_turn.rewrite, mem_rewrite)
                print(f"Memory Rewrite Retains First Entity: {'✅' if memory_rewrite_ok else '❌'}")
                
        except Exception as e:
            print(f"Error checking summary: {e}")
    
    return results


def print_summary(all_results: List[TestResult]):
    """Print overall test summary."""
    print(f"\n{'='*60}")
    print("OVERALL TEST SUMMARY")
    print(f"{'='*60}")
    
    by_conv = defaultdict(list)
    for r in all_results:
        by_conv[r.conversation_no].append(r)
    
    total_turns = len(all_results)
    total_rewrite_correct = sum(1 for r in all_results if r.rewrite_correct)
    total_answer_correct = sum(1 for r in all_results if r.answer_correct)
    
    print(f"\nTotal Turns Tested: {total_turns}")
    print(f"Rewrite Accuracy: {total_rewrite_correct}/{total_turns} ({total_rewrite_correct/total_turns*100:.1f}%)")
    print(f"Answer Accuracy: {total_answer_correct}/{total_turns} ({total_answer_correct/total_turns*100:.1f}%)")
    
    for conv_no, results in sorted(by_conv.items()):
        conv_rewrite = sum(1 for r in results if r.rewrite_correct)
        conv_answer = sum(1 for r in results if r.answer_correct)
        print(f"\nConversation {conv_no} ({len(results)} turns):")
        print(f"  Rewrite: {conv_rewrite}/{len(results)} ({conv_rewrite/len(results)*100:.1f}%)")
        print(f"  Answer: {conv_answer}/{len(results)} ({conv_answer/len(results)*100:.1f}%)")
        
        # Detail per turn
        for r in results:
            status = "✅" if (r.rewrite_correct and r.answer_correct) else "❌"
            print(f"  Turn {r.turn_no}: {status} Rewrite={r.rewrite_correct} Answer={r.answer_correct}")
    
    # Success criteria
    print(f"\n{'='*60}")
    print("SUCCESS CRITERIA CHECK")
    print(f"{'='*60}")
    
    rewrite_rate = total_rewrite_correct / total_turns if total_turns > 0 else 0
    answer_rate = total_answer_correct / total_turns if total_turns > 0 else 0
    
    criteria = [
        ("Rewrite > 70%", rewrite_rate > 0.7, f"Rewrite rate: {rewrite_rate*100:.1f}%"),
        ("Answer > 60%", answer_rate > 0.6, f"Answer rate: {answer_rate*100:.1f}%"),
    ]
    
    for name, passed, detail in criteria:
        print(f"  {'✅' if passed else '❌'} {name}: {detail}")
    
    # Memory retention check for conversations with 10 turns
    for conv_no, results in sorted(by_conv.items()):
        if len(results) >= 10:
            last_result = results[-1]
            try:
                detail = get_conversation_detail(last_result.conversation_id, f"QRECC_TEST_{conv_no}")
                context_summary = detail.get("context_summary", "").lower()
                first_rewrite = results[0].expected_rewrite.lower()
                first_words = set(w for w in first_rewrite.split() if len(w) > 4)
                retained = sum(1 for w in first_words if w in context_summary)
                memory_ok = retained >= len(first_words) * 0.5 if first_words else True
                print(f"  {'✅' if memory_ok else '❌'} Conversation {conv_no} Memory Retention: {retained}/{len(first_words)} key words in summary")
            except Exception:
                print(f"  ❌ Conversation {conv_no} Memory Retention: Could not check")


def main():
    print("="*60)
    print("QRECC MEMORY RAG TEST (TestClient)")
    print("="*60)
    
    # 1. Load data
    print("\n[1/5] Loading QRECC dataset...")
    turns = load_qrecc_data()
    print(f"Loaded {len(turns)} total turns")
    
    # 2. Group by conversation
    print("\n[2/5] Grouping by conversation...")
    grouped = group_by_conversation(turns)
    print(f"Found {len(grouped)} conversations")
    
    # 3. Select conversations
    print("\n[3/5] Selecting all conversations (>=2 turns)...")
    selected_convs = select_conversations(grouped, min_turns=2, max_convs=None)
    
    if not selected_convs:
        print("No conversations meet criteria!")
        return
    
    # 4. Create and upload documents
    print("\n[4/5] Creating and uploading documents...")
    doc_ids = {}
    test_dir = Path(__file__).resolve().parent / "qrecc_docs"
    test_dir.mkdir(exist_ok=True)
    
    for conv_no in selected_convs:
        conv_turns = grouped[conv_no]
        docx_path = test_dir / f"qrecc_conv_{conv_no}.docx"
        create_docx_for_conversation(conv_turns, docx_path)
        
        doc_id = upload_document(docx_path)
        if doc_id:
            if wait_for_processing(doc_id):
                doc_ids[conv_no] = doc_id
            else:
                print(f"  Skipping conversation {conv_no} - processing failed")
        else:
            print(f"  Skipping conversation {conv_no} - upload failed")
    
    if not doc_ids:
        print("No documents successfully processed!")
        return
    
    # 5. Run chat tests
    print("\n[5/5] Running chat queries...")
    all_results = []
    
    for conv_no in selected_convs:
        if conv_no not in doc_ids:
            continue
        
        user_id = f"QRECC_TEST_{conv_no}"
        conv_turns = grouped[conv_no]
        results = run_conversation_test(conv_no, conv_turns, user_id)
        all_results.extend(results)
    
    # Print summary
    print_summary(all_results)
    
    # Save detailed results
    output_file = test_dir / "qrecc_test_results.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump([
            {
                "conversation_no": r.conversation_no,
                "turn_no": r.turn_no,
                "question": r.question,
                "expected_rewrite": r.expected_rewrite,
                "actual_rewrite": r.actual_rewrite,
                "expected_answer": r.expected_answer,
                "actual_answer": r.actual_answer,
                "rewrite_correct": r.rewrite_correct,
                "answer_correct": r.answer_correct,
                "conversation_id": r.conversation_id
            }
            for r in all_results
        ], f, ensure_ascii=False, indent=2)
    
    print(f"\nDetailed results saved to: {output_file}")


if __name__ == "__main__":
    main()