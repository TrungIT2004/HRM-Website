from pathlib import Path
from typing import List
import openpyxl

from src.models.schema import DocumentBlock, ParsedDocument
from src.parsers.base_parser import BaseParser
from src.chunkers.table_preserver import TablePreserver


class ExcelParser(BaseParser):
    """
    Parser for Excel (.xlsx) files (salary tiers, KPI coefficients, leave matrices).
    Converts each sheet into structured markdown tables with sheet names as headers.
    """

    # Số dòng dữ liệu tối đa trước khi cắt thành nhiều bảng con (mỗi bảng con lặp lại header).
    # ponytail: ngưỡng này là cây chỉnh độ dài chunk, đổi được. Nhỏ quá -> mất ngữ cảnh dòng,
    # lớn quá -> vector bị loãng. 20 dòng x ~23 cột ~= 3-4 KB.
    MAX_DATA_ROWS_PER_SLICE = 20

    @staticmethod
    def _cell(val) -> str:
        """openpyxl trả số nguyên qua nhiều đường dưới dạng float; str() sẽ ra '21.0'."""
        if val is None:
            return ""
        if isinstance(val, float) and val.is_integer():
            return str(int(val))
        return str(val).strip()

    def parse(self, file_path: Path | str) -> ParsedDocument:
        path = Path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        blocks: List[DocumentBlock] = []

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            rows: List[List[str]] = []

            for row in sheet.iter_rows(values_only=True):
                str_row = [self._cell(val) for val in row]
                if any(val != "" for val in str_row):
                    rows.append(str_row)

            if rows:
                # Add sheet title as heading
                blocks.append(DocumentBlock(
                    block_type="HEADING",
                    content=f"Bảng dữ liệu: {sheet_name}",
                    heading_level=1
                ))

                # Bảng dài bị cắt thành nhiều TABLE block -> mỗi block thành 1 chunk riêng,
                # thay vì 1 chunk khống lồ cho cả sheet.
                for part_no, md_table in enumerate(
                    TablePreserver.to_markdown_slices(
                        rows,
                        caption=f"Bảng tra cứu - {sheet_name}",
                        max_data_rows=self.MAX_DATA_ROWS_PER_SLICE,
                    ),
                    start=1,
                ):
                    if not md_table:
                        continue
                    blocks.append(DocumentBlock(
                        block_type="TABLE",
                        content=md_table,
                        metadata={"sheet_name": sheet_name, "part": part_no}
                    ))

        wb.close()
        import re
        clean_name = re.sub(r'^[a-fA-F0-9]{8,64}_', '', path.name)
        return ParsedDocument(
            document_name=clean_name,
            file_type="xlsx",
            # ponytail: sheets are not pages, so page numbers stay unknown for xlsx.
            total_pages=0,
            blocks=blocks
        )
